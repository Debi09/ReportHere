"""ReportHere - FastAPI data quality engine (single file for Vercel)."""
import io
import math
import re
import warnings
from datetime import datetime, timezone

import numpy as np
import pandas as pd
from fastapi import FastAPI, File, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import JSONResponse, Response
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

warnings.filterwarnings("ignore")

app = FastAPI(title="ReportHere API")

ALLOWED_EXTENSIONS = {"csv", "xls", "xlsx"}
MAX_BYTES = 4 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 100
ID_COLUMNS = {"id", "order_id", "customer_id", "email", "employee_id"}
EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9\-]+(\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}$")

EMAIL_TOKENS = {"email", "emails", "mail"}
PHONE_TOKENS = {"phone", "mobile", "tel", "telephone", "cell", "fax"}
DATE_TOKENS = {"date", "dob", "birthday", "created", "updated", "timestamp", "joined", "hired"}
ID_TOKENS = {"id", "ids", "uuid", "guid"}
NUM_TOKENS = {"amount", "price", "salary", "cost", "qty", "quantity", "age", "total", "revenue",
              "score", "rate", "balance", "sales", "income", "weight", "height"}


# ---------------------------------------------------------------- errors
class ApiError(Exception):
    def __init__(self, status: int, message: str):
        self.status, self.message = status, message


@app.exception_handler(ApiError)
async def api_error_handler(_: Request, exc: ApiError):
    return JSONResponse(status_code=exc.status, content={"error": exc.message})


@app.exception_handler(RequestValidationError)
async def validation_handler(_: Request, __: RequestValidationError):
    return JSONResponse(status_code=400, content={"error": "No file received. Send a multipart field named 'file'."})


@app.exception_handler(Exception)
async def unexpected_handler(_: Request, __: Exception):
    return JSONResponse(status_code=500, content={"error": "Unexpected error while processing the file."})


# ---------------------------------------------------------------- small helpers
def tokens(name: str) -> set:
    spaced = re.sub(r"([a-z])([A-Z])", r"\1 \2", name)
    return set(re.split(r"[^a-z0-9]+", spaced.lower()))


def cell_text(v):
    if v is None or (not isinstance(v, (list, dict)) and pd.isna(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return str(v.date()) if v == v.normalize() else str(v)
    if isinstance(v, (float, np.floating)) and float(v).is_integer():
        return str(int(v))
    return str(v)


def jsafe(v):
    if v is None or v is pd.NaT:
        return None
    if isinstance(v, np.bool_):
        return bool(v)
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, (float, np.floating)):
        f = float(v)
        return None if (math.isnan(f) or math.isinf(f)) else round(f, 2)
    if isinstance(v, (bool, int, str)):
        return v
    return cell_text(v)


def to_text(vals: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(vals) and not pd.api.types.is_bool_dtype(vals):
        return vals.map(lambda x: str(int(x)) if float(x).is_integer() else str(x))
    return vals.astype(str).str.strip()


def to_num(vals: pd.Series) -> pd.Series:
    if pd.api.types.is_numeric_dtype(vals):
        return pd.to_numeric(vals, errors="coerce")
    cleaned = vals.astype(str).str.replace(r"[,$€£₹%\s]", "", regex=True)
    return pd.to_numeric(cleaned, errors="coerce")


def date_pattern(v: str) -> str:
    return re.sub(r"[A-Za-z]", "A", re.sub(r"\d", "9", v))


def examples_of(items, limit=10, quote=False):
    seen = list(dict.fromkeys(str(v) for v in items))[:limit]
    return [f'"{v}"' for v in seen] if quote else seen


def sanitize_excel_sheet_name(name) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", str(name)).strip().strip("'")
    return (cleaned or "Sheet")[:31]


# ---------------------------------------------------------------- reading
async def read_uploaded_file(file: UploadFile):
    filename = file.filename or "upload"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ApiError(400, "Unsupported file type. Please upload a CSV, XLS, or XLSX file.")

    content = await file.read(MAX_BYTES + 1)
    if not content:
        raise ApiError(400, "The uploaded file is empty.")
    if len(content) > MAX_BYTES:
        raise ApiError(413, "The file is larger than the 4 MB limit. Please upload a smaller file.")

    sheet = None
    df = None
    try:
        if ext == "csv":
            for enc in ("utf-8-sig", "latin-1"):
                try:
                    df = pd.read_csv(io.BytesIO(content), encoding=enc, skip_blank_lines=False)
                    break
                except UnicodeDecodeError:
                    continue
        else:
            engine = "openpyxl" if ext == "xlsx" else "xlrd"
            book = pd.ExcelFile(io.BytesIO(content), engine=engine)
            sheet = book.sheet_names[0]
            df = book.parse(sheet)
    except pd.errors.EmptyDataError:
        raise ApiError(422, "The file contains no data.")
    except Exception:
        raise ApiError(422, "Could not read this file. Make sure it is a valid, unprotected CSV or Excel file.")
    if df is None:
        raise ApiError(422, "Could not decode this file. Save it as UTF-8 CSV and try again.")

    df.columns = [str(c).strip() or f"column_{i + 1}" for i, c in enumerate(df.columns)]
    df = df.drop(columns=[c for c in df.columns if c.startswith("Unnamed:") and df[c].isna().all()])
    if df.shape[0] == 0 or df.shape[1] == 0:
        raise ApiError(422, "The file has no data rows or no columns.")
    if df.shape[0] > MAX_ROWS:
        raise ApiError(413, f"Too many rows ({df.shape[0]:,}). The limit is {MAX_ROWS:,} rows.")
    if df.shape[1] > MAX_COLUMNS:
        raise ApiError(413, f"Too many columns ({df.shape[1]}). The limit is {MAX_COLUMNS} columns.")

    info = {"filename": filename, "file_type": ext, "sheet_name": sheet,
            "file_size_kb": round(len(content) / 1024, 1)}
    return df, info


# ---------------------------------------------------------------- column analysis
def infer_type(col: str, vals: pd.Series) -> str:
    toks = tokens(col)
    if vals.empty:
        return "empty"
    if pd.api.types.is_bool_dtype(vals):
        return "boolean"
    if pd.api.types.is_datetime64_any_dtype(vals):
        return "date"
    numeric_dtype = pd.api.types.is_numeric_dtype(vals)
    text = to_text(vals.head(200))
    if toks & EMAIL_TOKENS or (not numeric_dtype and text.str.contains("@", regex=False).mean() >= 0.6):
        return "email"
    if toks & PHONE_TOKENS:
        return "phone"
    if toks & ID_TOKENS:
        return "identifier"
    if numeric_dtype:
        return "numeric"
    num_ratio = to_num(vals.head(200)).notna().mean()
    if num_ratio >= 0.8 or (toks & NUM_TOKENS and num_ratio >= 0.5):
        return "numeric"
    if num_ratio < 0.5:
        ratio = pd.to_datetime(text, errors="coerce", format="mixed").notna().mean()
        if ratio >= 0.8 or (toks & DATE_TOKENS and ratio >= 0.5):
            return "date"
    return "text"


def valid_phone(v: str) -> bool:
    digits = re.sub(r"\D", "", v)
    return bool(re.fullmatch(r"[0-9+\-().\sxX]+", v)) and 7 <= len(digits) <= 15


def column_status(missing_pct, n, issues) -> str:
    invalid = sum(i["count"] for i in issues if i["kind"] == "invalid")
    if missing_pct >= 30 or (n and invalid / n >= 0.10) or any(i["severity"] == "critical" for i in issues):
        return "Critical"
    if issues or missing_pct >= 5:
        return "Warning"
    return "Healthy"


def analyze_column(col, vals, blank_rows):
    """Return (inferred_type, issues, (min, max, mean))."""
    ctype = infer_type(col, vals)
    n = len(vals)
    issues = []
    stats = (None, None, None)

    def add(issue_type, kind, bad, quote=False):
        count = len(bad)
        if count == 0:
            return
        share = count / max(n, 1)
        severity = "critical" if kind != "format" and share >= 0.10 else "warning"
        issues.append({"column": col, "issue_type": issue_type, "kind": kind, "severity": severity,
                       "count": int(count), "examples": examples_of(bad, quote=quote)})

    add("Blank string", "format", blank_rows)

    if ctype == "email":
        t = to_text(vals)
        add("Invalid email format", "invalid", t[~t.map(lambda v: bool(EMAIL_RE.match(v)))])
    elif ctype == "phone":
        t = to_text(vals)
        add("Invalid phone number", "invalid", t[~t.map(valid_phone)])
    elif ctype == "identifier":
        t = to_text(vals)
        add("Duplicate ID", "duplicate", t[t.duplicated()])
    elif ctype == "date":
        if pd.api.types.is_datetime64_any_dtype(vals):
            parsed = vals
        else:
            t = to_text(vals)
            parsed = pd.to_datetime(t, errors="coerce", format="mixed")
            add("Unparseable date", "invalid", t[parsed.isna()])
            ok = t[parsed.notna()]
            pats = ok.map(date_pattern)
            if pats.nunique() > 1:
                add("Inconsistent date format", "format", ok[pats != pats.value_counts().index[0]])
        day = lambda ts: None if pd.isna(ts) else str(ts.date())
        stats = (day(parsed.min()), day(parsed.max()), None)
    elif ctype == "numeric":
        nums = to_num(vals)
        add("Non-numeric value", "invalid", vals[nums.isna()].astype(str))
        valid = nums.dropna()
        neg = valid[valid < 0]
        add("Negative value", "invalid", vals.loc[neg.index].astype(str))
        if len(valid) >= 8:
            q1, q3 = valid.quantile(0.25), valid.quantile(0.75)
            iqr = q3 - q1
            if iqr > 0:
                out = valid[(valid < q1 - 3 * iqr) | (valid > q3 + 3 * iqr)]
                out = out[~out.index.isin(neg.index)]
                add("Extreme outlier (IQR)", "invalid", vals.loc[out.index].astype(str))
        if len(valid):
            stats = (jsafe(valid.min()), jsafe(valid.max()), round(float(valid.mean()), 2))

    if ctype not in ("numeric", "date", "boolean", "empty"):
        sv = vals[vals.map(lambda x: isinstance(x, str))].astype(object)
        if len(sv):
            add("Leading/trailing whitespace", "format", sv[sv != sv.str.strip()], quote=True)

    if ctype == "text" and n:
        t = vals[vals.map(lambda x: isinstance(x, str))].astype(object)
        t = t.str.strip() if len(t) else t
        if 0 < t.nunique() <= max(20, 0.5 * len(t)):
            forms = {}
            for form, cnt in t.value_counts().items():
                forms.setdefault(form.lower(), []).append((form, cnt))
            bad_count, shown = 0, []
            for lst in forms.values():
                if len(lst) > 1:
                    bad_count += sum(c for _, c in lst[1:])
                    shown.append(" / ".join(f for f, _ in lst))
            if bad_count:
                issues.append({"column": col, "issue_type": "Inconsistent capitalization", "kind": "format",
                               "severity": "warning", "count": int(bad_count), "examples": shown[:10]})
    return ctype, issues, stats


# ---------------------------------------------------------------- main analysis
def analyze_dataframe(df, filename, info=None, sample_limit=10):
    info = info or {}
    df = df.copy()
    rows, cols = df.shape

    blank_mask = df.apply(lambda c: c.notna() & c.astype(str).str.strip().eq(""))
    work = df.mask(blank_mask)  # whitespace-only strings become missing
    empty_mask = work.isna().all(axis=1)
    empty_rows = int(empty_mask.sum())
    total_missing = int(work.isna().sum().sum())

    names = list(df.columns)
    dup_names = sorted({m.group(1) for n in names
                        if (m := re.match(r"^(.*)\.\d+$", n)) and m.group(1) in names})

    profiles, missing_list, issues = [], [], []
    for col in names:
        raw, vals = df[col], work[col].dropna()
        blank_rows = [f"Row {int(i) + 2}" for i in raw.index[blank_mask[col]]]
        ctype, col_issues, (mn, mx, mean) = analyze_column(col, vals, blank_rows)
        n = len(vals)
        missing = rows - n
        missing_pct = round(missing / rows * 100, 1)
        unique = int(vals.nunique())
        samples = list(dict.fromkeys(cell_text(v) for v in vals.head(50)))[:5]
        top = None
        if n:
            vc = vals.value_counts()
            if vc.iloc[0] > 1:
                top = cell_text(vc.index[0])
        profiles.append({
            "column": col, "inferred_type": ctype, "pandas_dtype": str(raw.dtype),
            "non_null_count": n, "missing_count": missing, "missing_pct": missing_pct,
            "unique_count": unique, "duplicate_value_count": n - unique,
            "min": mn, "max": mx, "mean": mean, "most_frequent": top,
            "invalid_count": sum(i["count"] for i in col_issues if i["kind"] == "invalid"),
            "issue_count": sum(i["count"] for i in col_issues),
            "status": column_status(missing_pct, n, col_issues),
            "sample_values": samples,
            "issues": [{k: i[k] for k in ("issue_type", "severity", "count", "examples")} for i in col_issues],
        })
        missing_list.append({"column": col, "missing_count": missing, "missing_pct": missing_pct,
                             "data_type": str(raw.dtype), "unique_count": unique, "examples": samples})
        issues.extend(col_issues)

    missing_list.sort(key=lambda m: (-m["missing_pct"], m["column"]))
    issues.sort(key=lambda i: (i["severity"] != "critical", -i["count"]))

    # duplicates
    body = work[~empty_mask]
    full_flags = body.duplicated(keep="first")
    full_dups = int(full_flags.sum())
    sources = []  # (match_type, ordered index list)
    if full_dups:
        grp = body[body.duplicated(keep=False)]
        try:
            order = grp.groupby(list(grp.columns), dropna=False, sort=False).ngroup().sort_values(kind="stable").index
        except Exception:
            order = grp.index
        sources.append(("Full row duplicate", list(order)))

    id_dups = []
    for col in names:
        if re.sub(r"[^a-z0-9]+", "_", col.lower()).strip("_") in ID_COLUMNS:
            t = to_text(work[col].dropna())
            if "email" in tokens(col):
                t = t.str.lower()
            flags = t.duplicated(keep="first")
            if flags.sum():
                id_dups.append({"column": col, "duplicate_rows": int(flags.sum()), "examples": examples_of(t[flags])})
                sources.append((f"Duplicate {col}", list(t[t.duplicated(keep=False)].sort_values().index)))

    samples, per_source = [], max(1, sample_limit // max(len(sources), 1))
    for match, idxs in sources:
        for idx in idxs[:per_source]:
            samples.append({"row_number": int(idx) + 2, "match_type": match,
                            "values": {c: cell_text(work.at[idx, c]) for c in names}})
    samples = samples[:sample_limit]

    invalid_count = sum(i["count"] for i in issues if i["kind"] == "invalid")
    format_count = sum(i["count"] for i in issues if i["kind"] == "format")
    dup_id_count = sum(i["count"] for i in issues if i["kind"] == "duplicate")

    results = {
        "file_metadata": {
            "filename": filename, "file_type": info.get("file_type", filename.rsplit(".", 1)[-1].lower()),
            "sheet_name": info.get("sheet_name"), "file_size_kb": info.get("file_size_kb", 0),
            "rows": rows, "columns": cols, "column_names": names, "duplicate_column_names": dup_names,
        },
        "summary": {
            "total_missing_values": total_missing, "duplicate_rows": full_dups, "empty_rows": empty_rows,
            "columns_with_issues": sum(1 for p in profiles if p["status"] != "Healthy"),
            "invalid_value_count": invalid_count, "format_issue_count": format_count,
            "duplicate_id_count": dup_id_count,
            "is_clean": not issues and total_missing == 0 and full_dups == 0 and not dup_names,
        },
        "column_profiles": profiles,
        "missing_values": missing_list,
        "validation_issues": issues,
        "duplicate_analysis": {"full_duplicate_rows": full_dups, "identifier_duplicates": id_dups,
                               "duplicate_column_names": dup_names},
        "duplicate_samples": samples,
    }
    results["quality_score"] = calculate_quality_score(results)
    results["insights"] = build_insights(results)
    results["report_generated_at"] = datetime.now(timezone.utc).isoformat()
    return results


def calculate_quality_score(results):
    """Start at 100; deduct capped points per factor. Never below 0."""
    m, s = results["file_metadata"], results["summary"]
    rows, cols = max(m["rows"], 1), max(m["columns"], 1)
    cells = rows * cols
    missing_cells = max(s["total_missing_values"] - s["empty_rows"] * cols, 0)  # blank rows scored separately
    factors = [
        ("Missing values", missing_cells / cells * 100, 1.0, 30),
        ("Duplicate rows", s["duplicate_rows"] / rows * 100, 1.5, 20),
        ("Invalid values", s["invalid_value_count"] / cells * 100, 3.0, 30),
        ("Blank rows", s["empty_rows"] / rows * 100, 1.0, 10),
    ]
    deductions = [{"factor": f, "percentage": round(p, 1), "points": round(min(cap, p * w), 1)}
                  for f, p, w, cap in factors]
    score = max(0, round(100 - sum(d["points"] for d in deductions)))
    grade, label = next((g, l) for t, g, l in
                        [(90, "A", "Excellent"), (75, "B", "Good"), (60, "C", "Needs Attention"),
                         (40, "D", "Poor"), (0, "F", "Critical")] if score >= t)
    parts = [f"{d['points']} for {d['factor'].lower()} ({d['percentage']}%)" for d in deductions if d["points"] > 0]
    explanation = ("Started at 100 and deducted " + ", ".join(parts) + "."
                   if parts else "No deductions: no missing values, duplicates, invalid values, or blank rows were found.")
    return {"score": score, "grade": grade, "label": label, "explanation": explanation, "deductions": deductions}


def build_insights(r):
    q, s, m = r["quality_score"], r["summary"], r["file_metadata"]
    profiles, issues = r["column_profiles"], r["validation_issues"]
    out = []
    sev = "success" if q["grade"] in ("A", "B") else "warning" if q["grade"] == "C" else "critical"
    out.append({"kind": "finding", "severity": sev,
                "text": f"The dataset scored {q['score']}/100 (grade {q['grade']}, {q['label']}) across {m['rows']:,} rows and {m['columns']} columns."})

    worst = max(profiles, key=lambda p: p["missing_count"] + p["issue_count"])
    if worst["missing_count"] + worst["issue_count"] > 0:
        top = max(worst["issues"], key=lambda i: i["count"])["issue_type"].lower() if worst["issues"] else None
        out.append({"kind": "finding", "severity": "warning",
                    "text": f"\"{worst['column']}\" is the most problematic column, with {worst['missing_count']} missing and {worst['issue_count']} flagged values" + (f" (mainly: {top})." if top else ".")})

    da = r["duplicate_analysis"]
    if s["duplicate_rows"] or da["identifier_duplicates"] or da["duplicate_column_names"]:
        bits = []
        if s["duplicate_rows"]:
            bits.append(f"{s['duplicate_rows']} fully duplicated row(s) ({s['duplicate_rows'] / max(m['rows'], 1) * 100:.1f}% of the data)")
        if da["identifier_duplicates"]:
            bits.append("repeated values in " + ", ".join(f"\"{d['column']}\"" for d in da["identifier_duplicates"]))
        if da["duplicate_column_names"]:
            bits.append("repeated column names (" + ", ".join(da["duplicate_column_names"]) + ")")
        out.append({"kind": "finding", "severity": "warning", "text": "Duplicates detected: " + "; ".join(bits) + "."})

    bad = [i for i in issues if i["kind"] == "invalid"]
    fmt = [i for i in issues if i["kind"] == "format"]
    if bad or fmt:
        parts = []
        if bad:
            b = max(bad, key=lambda i: i["count"])
            parts.append(f"{s['invalid_value_count']} invalid value(s), mostly \"{b['issue_type'].lower()}\" in \"{b['column']}\"")
        if fmt:
            f = max(fmt, key=lambda i: i["count"])
            parts.append(f"{s['format_issue_count']} formatting problem(s), mostly \"{f['issue_type'].lower()}\" in \"{f['column']}\"")
        out.append({"kind": "finding", "severity": "critical" if bad and any(i["severity"] == "critical" for i in bad) else "warning",
                    "text": "Found " + " and ".join(parts) + "."})

    top_factor = max(q["deductions"], key=lambda d: d["points"])
    if top_factor["points"] > 0:
        worst_missing = max(profiles, key=lambda p: p["missing_count"])["column"]
        action = {
            "Missing values": f"Fill or remove missing values, starting with \"{worst_missing}\".",
            "Duplicate rows": "Remove duplicate rows and enforce unique identifiers before using this data.",
            "Invalid values": "Correct the invalid entries listed in the Validation Issues sheet and add input validation at the source.",
            "Blank rows": f"Delete the {s['empty_rows']} blank row(s) before analysis.",
        }[top_factor["factor"]]
    elif fmt or s["duplicate_id_count"]:
        action = "Standardize whitespace, capitalization, date formats, and identifiers to make the data consistent."
    else:
        action = "No cleanup needed. The dataset is ready for analysis."
    out.append({"kind": "action", "severity": "info" if top_factor["points"] == 0 and not fmt else "warning",
                "text": "Recommended next action: " + action})
    return out[:5]


# ---------------------------------------------------------------- Excel report
NAVY, YELLOW, RED, GREEN = "0B1F3A", "FACC15", "E6392F", "86EFAC"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
FILLS = {"Healthy": PatternFill("solid", fgColor=GREEN), "Warning": PatternFill("solid", fgColor=YELLOW),
         "Critical": PatternFill("solid", fgColor=RED)}
_thin = Side(style="thin", color="9CA3AF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _xl(v):
    if v is None or (not isinstance(v, (list, dict)) and pd.isna(v)):
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().replace(tzinfo=None)
    if isinstance(v, (int, float, bool, datetime)):
        return v
    return ILLEGAL_CHARACTERS_RE.sub("", str(v))


def _write_table(ws, headers, rows, start=1, status_cols=(), filters=True):
    for j, h in enumerate(headers, 1):
        c = ws.cell(start, j, h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF", size=11), HEADER_FILL, BORDER
        c.alignment = Alignment(vertical="center", wrap_text=True)
    rows = list(rows)
    for i, row in enumerate(rows, start + 1):
        for j, raw in enumerate(row, 1):
            v = _xl(raw)
            c = ws.cell(i, j, v)
            if isinstance(v, str) and v.startswith("="):
                c.data_type = "s"  # never let uploaded text become a formula
            c.border, c.font = BORDER, Font(size=11)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if j in status_cols and v in FILLS:
                c.fill = FILLS[v]
                c.font = Font(size=11, bold=True, color="FFFFFF" if v == "Critical" else "111111")
    end = start + len(rows)
    if filters and rows:
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{end}"
    return end


def _finish(ws, widths=None):
    for cells in ws.columns:
        letter = get_column_letter(cells[0].column)
        longest = max((len(str(c.value)) for c in cells if c.value is not None), default=0)
        ws.column_dimensions[letter].width = (widths or {}).get(letter, min(max(longest + 2, 10), 50))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)


def generate_excel_report(results, dataframe) -> bytes:
    m, q, s = results["file_metadata"], results["quality_score"], results["summary"]
    wb = Workbook()

    # 1. Executive Summary
    ws = wb.active
    ws.title = sanitize_excel_sheet_name("Executive Summary")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "ReportHere: Data Quality Report"
    ws["A1"].font = Font(size=18, bold=True, color=NAVY)
    ws["A2"] = f"Generated: {results['report_generated_at']}"
    ws["A2"].font = Font(size=10, color="6B7280")
    row = _write_table(ws, ["File Information", "Value"], [
        ["File name", m["filename"]], ["File type", m["file_type"].upper()], ["Sheet analyzed", m["sheet_name"] or "n/a"],
        ["File size (KB)", m["file_size_kb"]], ["Rows", m["rows"]], ["Columns", m["columns"]]],
        start=4, filters=False) + 2
    score_row = row
    row = _write_table(ws, ["Data Quality Score", "Value"], [
        ["Score (0-100)", q["score"]], ["Grade", q["grade"]], ["Rating", q["label"]],
        ["How it was calculated", q["explanation"]]], start=row, filters=False) + 2
    fill = FILLS["Healthy"] if q["grade"] in ("A", "B") else FILLS["Warning"] if q["grade"] == "C" else FILLS["Critical"]
    for r_ in range(score_row + 1, score_row + 4):
        ws.cell(r_, 2).fill = fill
        ws.cell(r_, 2).font = Font(size=12, bold=True, color="FFFFFF" if q["grade"] in ("D", "F") else "111111")
    counts = [("Missing values", s["total_missing_values"]), ("Duplicate rows", s["duplicate_rows"]),
              ("Blank rows", s["empty_rows"]), ("Invalid values", s["invalid_value_count"]),
              ("Formatting issues", s["format_issue_count"]), ("Duplicate IDs", s["duplicate_id_count"]),
              ("Columns with issues", s["columns_with_issues"])]
    count_row = row
    row = _write_table(ws, ["Issue Counts", "Value"], counts, start=row, filters=False) + 2
    for i, (_, v) in enumerate(counts, count_row + 1):
        ws.cell(i, 2).fill = FILLS["Healthy"] if v == 0 else FILLS["Warning"]
    _write_table(ws, ["Type", "Automated Insight / Recommendation"],
                 [["Recommended action" if i["kind"] == "action" else "Finding", i["text"]] for i in results["insights"]],
                 start=row, filters=False)
    _finish(ws, {"A": 30, "B": 95})

    # 2. Column Profile
    ws = wb.create_sheet(sanitize_excel_sheet_name("Column Profile"))
    headers = ["Column", "Inferred Type", "Pandas Dtype", "Non-Null", "Missing", "Missing %", "Unique", "Duplicate Values",
               "Min", "Max", "Mean", "Most Frequent", "Invalid Values", "Status"]
    _write_table(ws, headers, [[p["column"], p["inferred_type"], p["pandas_dtype"], p["non_null_count"], p["missing_count"],
                                p["missing_pct"], p["unique_count"], p["duplicate_value_count"], p["min"], p["max"],
                                p["mean"], p["most_frequent"], p["invalid_count"], p["status"]]
                               for p in results["column_profiles"]], status_cols=(14,))
    ws.freeze_panes = "A2"
    _finish(ws)

    # 3. Missing Values
    ws = wb.create_sheet(sanitize_excel_sheet_name("Missing Values"))

    def miss_status(p):
        return "Healthy" if p == 0 else "Warning" if p < 30 else "Critical"

    _write_table(ws, ["Column", "Missing Count", "Missing %", "Data Type", "Unique Values", "Severity", "Example Values"],
                 [[v["column"], v["missing_count"], v["missing_pct"], v["data_type"], v["unique_count"],
                   miss_status(v["missing_pct"]), "; ".join(v["examples"])] for v in results["missing_values"]],
                 status_cols=(6,))
    ws.freeze_panes = "A2"
    _finish(ws)

    # 4. Validation Issues
    ws = wb.create_sheet(sanitize_excel_sheet_name("Validation Issues"))
    rows_ = [[i["issue_type"], i["column"], i["severity"].title(), i["count"], "; ".join(i["examples"])]
             for i in results["validation_issues"]] or [["No issues detected", "n/a", "Healthy", 0, ""]]
    _write_table(ws, ["Issue Type", "Affected Column", "Severity", "Issue Count", "Example Values"], rows_, status_cols=(3,))
    ws.freeze_panes = "A2"
    _finish(ws)

    # 5. Duplicate Records
    ws = wb.create_sheet(sanitize_excel_sheet_name("Duplicate Records"))
    da = results["duplicate_analysis"]
    summary_rows = [["Fully duplicated rows", da["full_duplicate_rows"]],
                    ["Duplicate column names", ", ".join(da["duplicate_column_names"]) or "None"]]
    summary_rows += [[f"Repeated values in '{d['column']}'", d["duplicate_rows"]] for d in da["identifier_duplicates"]]
    end = _write_table(ws, ["Duplicate Summary", "Value"], summary_rows, start=1, filters=False)
    cols = m["column_names"]
    samples = [[d["match_type"], d["row_number"]] + [d["values"].get(c) for c in cols] for d in results["duplicate_samples"]]
    if samples:
        _write_table(ws, ["Match Type", "Row Number"] + cols, samples, start=end + 2)
    else:
        ws.cell(end + 2, 1, "No duplicate records detected.").font = Font(bold=True)
    ws.freeze_panes = "A2"
    _finish(ws)

    # 6. Cleaned Data Preview
    ws = wb.create_sheet(sanitize_excel_sheet_name("Cleaned Data Preview"))
    clean = dataframe.copy()
    for c in clean.columns:
        if not pd.api.types.is_numeric_dtype(clean[c]) and not pd.api.types.is_datetime64_any_dtype(clean[c]):
            clean[c] = clean[c].map(lambda v: v.strip() if isinstance(v, str) else v)
            if "email" in tokens(c):
                clean[c] = clean[c].map(lambda v: v.lower() if isinstance(v, str) else v)
    clean = clean.replace(r"^\s*$", np.nan, regex=True).dropna(how="all").drop_duplicates().head(500)
    _write_table(ws, list(clean.columns), clean.itertuples(index=False, name=None))
    ws.freeze_panes = "A2"
    _finish(ws)

    for sheet in wb.worksheets[1:]:
        sheet.print_title_rows = "1:1"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- endpoints
@app.get("/api/health")
def health():
    return {"status": "ok", "service": "data-quality-reporting-api"}


@app.post("/api/analyze")
async def analyze(file: UploadFile = File(...)):
    df, info = await read_uploaded_file(file)
    return analyze_dataframe(df, info["filename"], info)


@app.post("/api/generate-report")
async def generate_report(file: UploadFile = File(...)):
    df, info = await read_uploaded_file(file)
    results = analyze_dataframe(df, info["filename"], info)
    data = generate_excel_report(results, df)
    name = f"data_quality_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )
